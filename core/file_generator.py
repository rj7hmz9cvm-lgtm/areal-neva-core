from pathlib import Path

def save_txt(path, text):
    Path(path).write_text(text or "", encoding="utf-8")

def save_docx(path, text):
    from docx import Document
    doc = Document()
    for line in (text or "").splitlines():
        doc.add_paragraph(line)
    doc.save(path)

def save_xlsx(path, text):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    
    row_idx = 1
    for line in (text or "").splitlines():
        line = line.strip()
        if not line or set(line) <= {"-", "|", " "}: # Пропускаем пустые и разделители таблиц
            continue
            
        # Убираем крайние палки Markdown таблицы
        if line.startswith("|"): line = line[1:]
        if line.endswith("|"): line = line[:-1]
        
        cols = [c.strip() for c in line.split("|")]
        # Если это реально табличная строка
        if len(cols) > 1:
            for col_idx, val in enumerate(cols, 1):
                ws.cell(row=row_idx, column=col_idx, value=val)
            row_idx += 1

    wb.save(path)
